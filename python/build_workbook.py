import pandas as pd, datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

raw = pd.read_csv("data/Raw_Positions.csv", dtype={"AssetClass": str}, keep_default_na=False)
thr = pd.read_csv("data/Thresholds.csv")
N = len(raw); LAST = N + 1
F = "Arial"
base = Font(name=F, size=10); bold = Font(name=F, size=10, bold=True)
hdr_font = Font(name=F, size=10, bold=True, color="FFFFFF"); hdr_fill = PatternFill("solid", fgColor="1F3864")
blue = Font(name=F, size=10, color="0000FF"); green = Font(name=F, size=10, color="008000")
yellow = PatternFill("solid", fgColor="FFFF00")
thin = Side(style="thin", color="BFBFBF")

def header(ws, row, labels, col=1):
    for i, l in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=l); c.font = hdr_font; c.fill = hdr_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

wb = Workbook()
# ---------------- README ----------------
ws = wb.active; ws.title = "README"
lines = [
 ("Position Price Verification - Valuation Control Workbook", "title"),
 ("SYNTHETIC DATA ONLY. No real bank, client or market data is used. Built as a learning / portfolio project.", "warn"),
 ("", None),
 ("Purpose", "h"),
 ("Compare desk marks (BookPrice) with independent prices, flag positions outside asset-class tolerance, and report data-quality issues.", None),
 ("", None),
 ("Sheets", "h"),
 ("Thresholds - tolerance per asset class and the Report Date (inputs, blue font).", None),
 ("Raw_Data - deliberately messy input: duplicates, missing independent prices, stale dates, text-case/space issues, one unmapped asset class label.", None),
 ("Checks - formula-driven control test on every Raw_Data row (works without macros).", None),
 ("Summary - breach counts, rates and net market-value differences by asset class and by desk, plus data-quality counts.", None),
 ("VBA (separate file ValuationAutomation.bas) - one-click automation: cleans, tests, builds Clean_Data + Exceptions + Run_Log, exports Clean_Positions.csv.", None),
 ("", None),
 ("Rules", "h"),
 ("Variance % = (BookPrice - IndependentPrice) / IndependentPrice.", None),
 ("Status priority: UNMAPPED ASSET CLASS > NO INDEPENDENT PRICE > BREACH (|Variance %| > tolerance) > OK.", None),
 ("Market value difference = Quantity x (BookPrice - IndependentPrice), in local currency, price per unit. No FX conversion (assumption).", None),
 ("Duplicates: first occurrence of a PositionID is kept; later copies are flagged DUP and excluded from Summary.", None),
 ("Breach rate = Breaches / (OK + Breaches), i.e. positions that could actually be tested.", None),
 ("", None),
 ("Assumptions", "h"),
 ("Tolerances (Equity 1.00%, Corporate Bond 0.50%, Government Bond 0.25%, FX Forward 0.30%) are illustrative values chosen by the author, not any bank's policy.", None),
 ("Report Date 28-Sep-2026 is an input; rows with a different ValuationDate are flagged STALE.", None),
 ("", None),
 ("Colour legend", "h"),
 ("Blue font = hard-coded input | Black = formula | Green = link to another sheet | Yellow fill = key input to change.", None),
]
for i, (txt, kind) in enumerate(lines, 1):
    c = ws.cell(row=i, column=1, value=txt); c.font = base; c.alignment = Alignment(wrap_text=True, vertical="top")
    if kind == "title": c.font = Font(name=F, size=14, bold=True, color="1F3864")
    if kind == "h": c.font = Font(name=F, size=11, bold=True, color="1F3864")
    if kind == "warn": c.font = Font(name=F, size=10, bold=True, color="C00000")
ws.column_dimensions["A"].width = 120

# ---------------- Thresholds ----------------
wt = wb.create_sheet("Thresholds")
header(wt, 1, ["Asset Class", "Tolerance (%)"])
for i, r in enumerate(thr.itertuples(), 2):
    wt.cell(row=i, column=1, value=r.AssetClass).font = blue
    c = wt.cell(row=i, column=2, value=float(r.TolerancePct)); c.font = blue; c.number_format = "0.00%"
wt["D1"] = "Report Date"; wt["D1"].font = bold
wt["E1"] = dt.date(2026, 9, 28); wt["E1"].font = blue; wt["E1"].fill = yellow; wt["E1"].number_format = "yyyy-mm-dd"
wt["A7"] = "Note: tolerances are illustrative author assumptions. Edit blue cells to test other thresholds; Checks and Summary recalculate."; wt["A7"].font = Font(name=F, size=9, italic=True)
wt.column_dimensions["A"].width = 20; wt.column_dimensions["B"].width = 14; wt.column_dimensions["D"].width = 14; wt.column_dimensions["E"].width = 14

# ---------------- Raw_Data ----------------
wr = wb.create_sheet("Raw_Data")
cols = list(raw.columns)
header(wr, 1, cols)
for i, r in enumerate(raw.itertuples(index=False), 2):
    for j, v in enumerate(r, 1):
        name = cols[j - 1]
        if name == "ValuationDate":
            v = dt.datetime.strptime(v, "%Y-%m-%d").date()
        elif name in ("Quantity",):
            v = int(v)
        elif name in ("BookPrice", "IndependentPrice"):
            v = None if v == "" else float(v)
        c = wr.cell(row=i, column=j, value=v); c.font = base
        if name == "ValuationDate": c.number_format = "yyyy-mm-dd"
        if name == "Quantity": c.number_format = "#,##0"
        if name in ("BookPrice", "IndependentPrice"): c.number_format = "#,##0.00000"
for j, w in enumerate([11, 13, 20, 18, 12, 9, 13, 13, 16], 1): wr.column_dimensions[get_column_letter(j)].width = w
wr.freeze_panes = "A2"

# ---------------- Checks ----------------
wc = wb.create_sheet("Checks")
header(wc, 1, ["PositionID", "Desk", "Asset Class (clean)", "Quantity", "Book Price", "Independent Price", "Variance %",
               "Tolerance %", "MV Diff (local ccy)", "Status", "Duplicate?", "Valuation Date Flag"])
wc.row_dimensions[1].height = 30
TA, TB = "Thresholds!$A$2:$A$5", "Thresholds!$B$2:$B$5"
for r in range(2, LAST + 1):
    f = {
     "A": f"=Raw_Data!A{r}", "B": f"=Raw_Data!C{r}",
     "C": f'=IFERROR(INDEX({TA},MATCH(TRIM(Raw_Data!D{r}),{TA},0)),"UNMAPPED")',
     "D": f"=Raw_Data!G{r}", "E": f"=Raw_Data!H{r}",
     "F": f'=IF(Raw_Data!I{r}="","",Raw_Data!I{r})',
     "G": f'=IF(F{r}="","",(E{r}-F{r})/F{r})',
     "H": f'=IFERROR(INDEX({TB},MATCH(C{r},{TA},0)),"")',
     "I": f'=IF(F{r}="","",D{r}*(E{r}-F{r}))',
     "J": f'=IF(C{r}="UNMAPPED","UNMAPPED ASSET CLASS",IF(F{r}="","NO INDEPENDENT PRICE",IF(ABS(G{r})>H{r},"BREACH","OK")))',
     "K": f'=IF(COUNTIF(Raw_Data!$A$2:A{r},Raw_Data!A{r})>1,"DUP","UNIQUE")',
     "L": f'=IF(Raw_Data!B{r}<>Thresholds!$E$1,"STALE","CURRENT")',
    }
    for col, formula in f.items():
        c = wc[f"{col}{r}"]; c.value = formula
        c.font = green if col in "ABDE" else base
    wc[f"D{r}"].number_format = "#,##0"; wc[f"E{r}"].number_format = "#,##0.00000"; wc[f"F{r}"].number_format = "#,##0.00000"
    wc[f"G{r}"].number_format = "0.000%"; wc[f"H{r}"].number_format = "0.00%"; wc[f"I{r}"].number_format = "#,##0.00;(#,##0.00);-"
for j, w in enumerate([11, 20, 18, 12, 13, 16, 12, 12, 18, 24, 12, 16], 1): wc.column_dimensions[get_column_letter(j)].width = w
wc.freeze_panes = "B2"; wc.auto_filter.ref = f"A1:L{LAST}"
rng = f"J2:J{LAST}"
wc.conditional_formatting.add(rng, FormulaRule(formula=['J2="BREACH"'], fill=PatternFill("solid", bgColor="F8CBAD", fgColor="F8CBAD")))
wc.conditional_formatting.add(rng, FormulaRule(formula=['OR(J2="NO INDEPENDENT PRICE",J2="UNMAPPED ASSET CLASS")'], fill=PatternFill("solid", bgColor="FFE699", fgColor="FFE699")))
wc.conditional_formatting.add(f"K2:K{LAST}", FormulaRule(formula=['K2="DUP"'], fill=PatternFill("solid", bgColor="D9D9D9", fgColor="D9D9D9")))
wc.conditional_formatting.add(f"L2:L{LAST}", FormulaRule(formula=['L2="STALE"'], fill=PatternFill("solid", bgColor="FFE699", fgColor="FFE699")))

# ---------------- Summary ----------------
wsu = wb.create_sheet("Summary")
def rngc(col): return f"Checks!${col}$2:${col}${LAST}"
wsu["A1"] = "Valuation Control Summary"; wsu["A1"].font = Font(name=F, size=14, bold=True, color="1F3864")
wsu["A2"] = "Report Date"; wsu["A2"].font = bold; wsu["B2"] = "=Thresholds!E1"; wsu["B2"].font = green; wsu["B2"].number_format = "yyyy-mm-dd"
labels = ["Positions Tested", "OK", "Breaches", "No Indep. Price", "Breach Rate", "Net MV Diff (all tested)", "Net MV Diff (breaches)"]
def block(start, title, keycol, keys, key_is_formula):
    wsu.cell(row=start, column=1, value=title).font = Font(name=F, size=11, bold=True, color="1F3864")
    header(wsu, start + 1, [title.split(" by ")[-1].title()] + labels)
    wsu.row_dimensions[start + 1].height = 30
    r0 = start + 2
    for k, key in enumerate(keys):
        r = r0 + k
        a = wsu.cell(row=r, column=1, value=(f"=Thresholds!A{k+2}" if key_is_formula else key)); a.font = green if key_is_formula else base
        crit = f"{rngc(keycol)},$A{r},{rngc('K')},\"UNIQUE\""
        wsu[f"B{r}"] = f'=COUNTIFS({rngc(keycol)},$A{r},{rngc("K")},"UNIQUE",{rngc("J")},"OK")+COUNTIFS({crit},{rngc("J")},"BREACH")'
        wsu[f"C{r}"] = f'=COUNTIFS({crit},{rngc("J")},"OK")'
        wsu[f"D{r}"] = f'=COUNTIFS({crit},{rngc("J")},"BREACH")'
        wsu[f"E{r}"] = f'=COUNTIFS({crit},{rngc("J")},"NO INDEPENDENT PRICE")'
        wsu[f"F{r}"] = f"=IF((C{r}+D{r})=0,0,D{r}/(C{r}+D{r}))"
        wsu[f"G{r}"] = f'=SUMIFS({rngc("I")},{crit},{rngc("J")},"<>UNMAPPED ASSET CLASS")'
        wsu[f"H{r}"] = f'=SUMIFS({rngc("I")},{crit},{rngc("J")},"BREACH")'
    rt = r0 + len(keys)
    wsu.cell(row=rt, column=1, value="Total").font = bold
    for col in "BCDE": wsu[f"{col}{rt}"] = f"=SUM({col}{r0}:{col}{rt-1})"
    wsu[f"F{rt}"] = f"=IF((C{rt}+D{rt})=0,0,D{rt}/(C{rt}+D{rt}))"
    for col in "GH": wsu[f"{col}{rt}"] = f"=SUM({col}{r0}:{col}{rt-1})"
    for r in range(r0, rt + 1):
        for col in "BCDEFGH":
            c = wsu[f"{col}{r}"]; c.font = bold if r == rt else base
            c.number_format = {"F": "0.0%", "G": "#,##0;(#,##0);-", "H": "#,##0;(#,##0);-"}.get(col, "#,##0")
        if r == rt:
            for col in "ABCDEFGH": wsu[f"{col}{r}"].border = Border(top=thin)
    return rt
end1 = block(4, "Results by Asset Class", "C", list(thr.AssetClass), True)
desks = sorted(raw["Desk"].unique())
end2 = block(end1 + 3, "Results by Desk", "B", desks, False)
q = end2 + 3
wsu.cell(row=q, column=1, value="Data Quality").font = Font(name=F, size=11, bold=True, color="1F3864")
header(wsu, q + 1, ["Check", "Count"])
dq = [("Raw rows loaded", "=COUNTA(Raw_Data!A2:A%d)" % LAST),
      ("Duplicate rows removed", f'=COUNTIF({rngc("K")},"DUP")'),
      ("Unique positions", f'=COUNTIF({rngc("K")},"UNIQUE")'),
      ("Missing independent price", f'=COUNTIFS({rngc("J")},"NO INDEPENDENT PRICE",{rngc("K")},"UNIQUE")'),
      ("Unmapped asset class", f'=COUNTIFS({rngc("J")},"UNMAPPED ASSET CLASS",{rngc("K")},"UNIQUE")'),
      ("Stale valuation date", f'=COUNTIFS({rngc("L")},"STALE",{rngc("K")},"UNIQUE")')]
for i, (l, fm) in enumerate(dq):
    wsu.cell(row=q + 2 + i, column=1, value=l).font = base
    c = wsu.cell(row=q + 2 + i, column=2, value=fm); c.font = base; c.number_format = "#,##0"
wsu.column_dimensions["A"].width = 26
for col in "BCDEFGH": wsu.column_dimensions[col].width = 16
wb.save("excel/Valuation_Control_Workbook.xlsx")
print("saved", LAST, "rows; summary rows:", end1, end2, q)
